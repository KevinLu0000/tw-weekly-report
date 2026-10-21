# 產生每週的資料試算表：results/<W>/台股週報資料_<W>.xlsx
# 工作表：價值股Top10、價值股第一階段候選、動量股Top10、動量股全部通過、產業強弱、風險計分
# 數字都是程式算好的結果（不是公式），每張表第一列寫資料日期與來源
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="DDE7F3")       # 表頭底色（淡藍）


def sheet(wb, title, note, cols, rows):
    """建立一張工作表：第 1 列說明、第 2 列表頭、第 3 列起資料
    cols：[(表頭, 取值函式, 數字格式), ...]"""
    ws = wb.create_sheet(title)
    ws.append([note])
    ws["A1"].font = Font(italic=True, color="666666")
    ws.append([c[0] for c in cols])
    for cell in ws[2]:
        cell.font, cell.fill = Font(bold=True), HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r in rows:
        vals = []
        for _, fn, _ in cols:
            try:
                vals.append(fn(r))
            except (KeyError, TypeError, IndexError):
                vals.append(None)                    # 缺資料就留空白
        ws.append(vals)
    for j, (_, _, fmt) in enumerate(cols, 1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = max(9, min(28, len(cols[j - 1][0]) * 2 + 2))
        if fmt:
            for i in range(3, ws.max_row + 1):
                ws[f"{letter}{i}"].number_format = fmt
    ws.freeze_panes = "C3"
    return ws


def build(out, W, v_top, v_backup, dropped, v_all, m_top, m_all, mk):
    wb = Workbook()
    wb.remove(wb.active)
    src = f"資料截至 {W[:4]}/{W[4:6]}/{W[6:]} 收盤；來源：證交所、櫃買中心 OpenAPI 與每日行情、Yahoo 股市、HiStock；規則見 RULES.md"
    pct, d1, d2, n0 = '0.0"%"', "0.0", "0.00", "#,##0"
    vcols = [("代號", lambda s: s["code"], None), ("公司", lambda s: s["name"], None), ("市場", lambda s: s["mkt"], None),
             ("產業", lambda s: s.get("ind"), None), ("價值評分", lambda s: s["score"], d1),
             ("市值(億)", lambda s: s["mcap"], n0), ("收盤", lambda s: s["close"], d2),
             ("52週低", lambda s: s["low52"], d2), ("距低點%", lambda s: s["dist"], pct),
             ("52週高", lambda s: s["high52"], d2), ("距高點%", lambda s: s["from_hi"], pct),
             ("本益比", lambda s: s["pe"], d1), ("淨值比", lambda s: s["pb"], d2), ("殖利率%", lambda s: s["dy"], pct),
             ("最新月營收年增%", lambda s: s.get("rev_yoy"), pct), ("累計營收年增%", lambda s: s.get("rev_cum_yoy"), pct),
             ("負債比%", lambda s: s.get("debt"), pct), ("ROE%", lambda s: s.get("roe"), pct),
             ("近四季EPS", lambda s: s.get("eps4"), d2), ("近四季營業CF(億)", lambda s: s.get("ocf4"), d1),
             ("近四季自由CF(億)", lambda s: s.get("fcf4"), d1), ("融資5日增減%", lambda s: s.get("融資5日增減%"), pct),
             ("券資比%", lambda s: s.get("券資比%"), pct)]
    sheet(wb, "價值股Top10", src, [("名次", lambda s: s["_rank"], None)] + vcols,
          [dict(s, _rank=i) for i, s in enumerate(v_top, 1)] + [dict(s, _rank="候補") for s in v_backup])
    status = {s["code"]: "前10名" for s in v_top} | {s["code"]: "候補" for s in v_backup} | \
             {s["code"]: "剔除：" + s["drop_reason"] for s in dropped}
    sheet(wb, "價值股第一階段候選", src + "；第二階段只檢查到湊滿 11 檔為止",
          vcols[:2] + [("第二階段結果", lambda s: status.get(s["code"], "未檢查（排名較後）"), None)] + vcols[2:18],
          v_all)
    mcols = [("代號", lambda s: s["code"], None), ("公司", lambda s: s["name"], None), ("市場", lambda s: s["mkt"], None),
             ("產業", lambda s: s.get("ind"), None), ("總分", lambda s: s["score"], d1),
             ("1週%", lambda s: s["r5"], pct), ("1月%", lambda s: s["r22"], pct), ("3月%", lambda s: s["r65"], pct),
             ("52週高", lambda s: s["hi52"], d2), ("距高點%", lambda s: s["from_hi"], pct),
             ("52週低", lambda s: s["lo52"], d2), ("距低點%", lambda s: s["from_lo"], pct),
             ("外資5日(張)", lambda s: s["f5"], n0), ("投信5日(張)", lambda s: s["t5"], n0),
             ("法人買超佔成交值%", lambda s: s["inst_ratio"], pct), ("22日線", lambda s: s["ma22"], d1),
             ("65日線", lambda s: s["ma65"], d1), ("20日均成交值(億)", lambda s: s["val20"], d1),
             ("處置股", lambda s: "是" if s["punish"] else "", None), ("注意股", lambda s: "是" if s["notice"] else "", None)]
    sheet(wb, "動量股Top10", src, [("名次", lambda s: s["_rank"], None)] + mcols[:4] +
          [("分類", lambda s: s.get("class"), None), ("標示", lambda s: "、".join(s.get("tags", [])), None)] + mcols[4:] +
          [("融資5日增減%", lambda s: s.get("融資5日增減%"), pct), ("券資比%", lambda s: s.get("券資比%"), pct)],
          [dict(s, _rank=i) for i, s in enumerate(m_top, 1)])
    sheet(wb, "動量股全部通過", src, mcols, m_all)
    s4 = mk["產業"]
    sheet(wb, "產業強弱", src + f"；等權重；營收月份 {s4.get('營收月份')}",
          [("產業", lambda x: x["產業"], None), ("檔數", lambda x: x["檔數"], None), ("上漲", lambda x: x["上漲"], None),
           ("週報酬%", lambda x: x["週報酬%"], pct), ("週中位數%", lambda x: x["週中位數%"], pct),
           ("1月報酬%", lambda x: x["1月報酬%"], pct), ("營收年增中位數%", lambda x: x["營收年增中位數%"], pct),
           ("營收年增家數", lambda x: x["營收年增家數"], None), ("營收家數", lambda x: x["營收家數"], None)],
          s4["產業"])
    r = mk["風險"]
    ws = sheet(wb, "風險計分", src + f"；總分 {r['總分']}，判斷：{r['判斷']}",
               [("訊號", lambda x: x["訊號"], None), ("本週數字", lambda x: str(x["數字"]), None),
                ("得分", lambda x: x["得分"], None)], r["訊號"])
    ws.append(["合計", "", r["總分"]])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 36
    wb.save(f"{out}/台股週報資料_{W}.xlsx")
